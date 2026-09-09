#!/usr/bin/env python3
"""Đối chiếu nhãn của nhóm khác với ground truth cục bộ và xuất Excel.

Script không sửa workbook đầu vào. Mỗi record được ghép nghiêm ngặt bằng tên
ảnh, không ghép theo số thứ tự hoặc nội dung gần đúng.
"""

from __future__ import annotations

import argparse
import difflib
import json
import os
from copy import copy
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "Nhóm 2_validation_sent.xlsx"
DEFAULT_GROUND_TRUTH = ROOT / "data" / "output" / "DeepSeek_ground_truth" / "ground_truth.jsonl"
DEFAULT_TEMPLATE = ROOT / "Kết quả đánh giá chéo 1-100.xlsx"
DEFAULT_OUTPUT = ROOT / "Kết quả đánh giá chéo 201-350.xlsx"

OUTPUT_HEADERS = (
    "No",
    "image",
    "ground_truth",
    "label",
    "Corrected",
    "Levenshtein Accuracy",
    "Levenshtein Accuracy (bao gồm cả line break)",
    "Note",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--ground-truth", type=Path, default=DEFAULT_GROUND_TRUTH)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--start-id", type=int, default=201)
    parser.add_argument("--end-id", type=int, default=350)
    args = parser.parse_args()
    if args.start_id < 1 or args.end_id < args.start_id:
        parser.error("Cần 1 <= start-id <= end-id")
    if args.output.resolve() in {args.input.resolve(), args.template.resolve()}:
        parser.error("Output phải khác input và template để không ghi đè file gốc")
    return args


def read_jsonl(path: Path) -> Iterable[dict[str, object]]:
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"JSON lỗi tại {path}:{line_number}: {exc}") from exc
            if not isinstance(row, dict):
                raise ValueError(f"Dòng {line_number} trong {path} không phải object")
            yield row


def load_ground_truth(path: Path) -> dict[str, str]:
    indexed: dict[str, str] = {}
    for row in read_jsonl(path):
        image = Path(str(row.get("image", ""))).name
        text = str(row.get("ground_truth", ""))
        if not image:
            continue
        if image in indexed:
            raise ValueError(f"Tên ảnh bị trùng trong ground truth: {image}")
        indexed[image] = text
    if not indexed:
        raise ValueError(f"Ground truth rỗng: {path}")
    return indexed


def header_map(worksheet) -> dict[str, int]:
    result = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    required = {"#", "Image", "FB Caption", "Label"}
    missing = sorted(required - set(result))
    if missing:
        raise ValueError(f"Workbook input thiếu cột: {', '.join(missing)}")
    return result


def levenshtein(left: str, right: str) -> int:
    if len(left) < len(right):
        left, right = right, left
    previous = list(range(len(right) + 1))
    for left_index, left_char in enumerate(left, 1):
        current = [left_index]
        for right_index, right_char in enumerate(right, 1):
            current.append(
                min(
                    current[-1] + 1,
                    previous[right_index] + 1,
                    previous[right_index - 1] + (left_char != right_char),
                )
            )
        previous = current
    return previous[-1]


def accuracy(reference: str, candidate: str, include_line_breaks: bool) -> float:
    reference = reference.replace("\r\n", "\n").replace("\r", "\n")
    candidate = candidate.replace("\r\n", "\n").replace("\r", "\n")
    if not include_line_breaks:
        reference = reference.replace("\n", "")
        candidate = candidate.replace("\n", "")
    denominator = max(len(reference), len(candidate))
    if denominator == 0:
        return 1.0
    return max(0.0, 1.0 - levenshtein(reference, candidate) / denominator)


def normalized_newlines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def layout_signature(text: str) -> tuple[str, set[int], set[int]]:
    """Return non-whitespace content and line/space positions in that content."""
    content: list[str] = []
    line_breaks: set[int] = set()
    spaces: set[int] = set()
    for char in normalized_newlines(text):
        if char == "\n":
            line_breaks.add(len(content))
        elif char.isspace():
            spaces.add(len(content))
        else:
            content.append(char)
    return "".join(content), line_breaks, spaces


def boundary_context(content: str, position: int, radius: int = 6) -> str:
    left = content[max(0, position - radius):position]
    right = content[position:position + radius]
    return f"“{left}｜{right}”"


def summarize_boundaries(content: str, positions: set[int], limit: int = 4) -> str:
    ordered = sorted(positions)
    examples = [boundary_context(content, position) for position in ordered[:limit]]
    suffix = f" và {len(ordered) - limit} vị trí khác" if len(ordered) > limit else ""
    return ", ".join(examples) + suffix


def textual_edits(group2_text: str, group1_text: str, limit: int = 6) -> list[str]:
    group2 = "".join(group2_text.split())
    group1 = "".join(group1_text.split())
    edits: list[str] = []
    matcher = difflib.SequenceMatcher(None, group1, group2, autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        expected = group1[i1:i2]
        observed = group2[j1:j2]
        if tag == "replace":
            edits.append(f"Nhóm 2 ghi “{observed}” thay cho “{expected}”")
        elif tag == "delete":
            edits.append(f"Nhóm 2 thiếu “{expected}”")
        elif tag == "insert":
            edits.append(f"Nhóm 2 thừa “{observed}”")
        if len(edits) == limit:
            edits.append("còn sai khác khác")
            break
    return edits


def note_for(group2_text: str, group1_text: str, score: float, line_score: float) -> str:
    if group2_text == group1_text:
        return "Khớp hoàn toàn với ground truth của Nhóm 1."

    group2_content, group2_breaks, group2_spaces = layout_signature(group2_text)
    group1_content, group1_breaks, group1_spaces = layout_signature(group1_text)
    if group2_content == group1_content:
        group2_lines = len(normalized_newlines(group2_text).splitlines())
        group1_lines = len(normalized_newlines(group1_text).splitlines())
        missing_breaks = group1_breaks - group2_breaks
        extra_breaks = group2_breaks - group1_breaks
        details = [
            f"Toàn bộ {len(group1_content)} ký tự nội dung khớp ground truth Nhóm 1; "
            "không có lỗi nhận dạng chữ."
        ]
        if group2_lines != group1_lines:
            details.append(
                f"Bố cục khác: Nhóm 2 trình bày {group2_lines} dòng, "
                f"ground truth Nhóm 1 có {group1_lines} dòng."
            )
        if missing_breaks:
            details.append(
                f"Nhóm 2 đã gộp/thiếu {len(missing_breaks)} điểm xuống dòng, tại "
                f"{summarize_boundaries(group1_content, missing_breaks)}."
            )
        if extra_breaks:
            details.append(
                f"Nhóm 2 đặt thêm {len(extra_breaks)} điểm xuống dòng, tại "
                f"{summarize_boundaries(group1_content, extra_breaks)}."
            )
        if group2_spaces != group1_spaces:
            details.append(
                f"Khoảng trắng cũng khác ({len(group2_spaces)} vị trí ở Nhóm 2, "
                f"{len(group1_spaces)} vị trí ở ground truth Nhóm 1)."
            )
        details.append(
            f"Accuracy bỏ xuống dòng={score:.2%}; giữ xuống dòng={line_score:.2%}. "
            "Chỉ cần chuẩn hóa cách ngắt dòng/khoảng trắng theo cột Corrected."
        )
        return " ".join(details)

    distance = levenshtein(group1_text.replace("\n", ""), group2_text.replace("\n", ""))
    edits = textual_edits(group2_text, group1_text)
    edit_summary = "; ".join(edits) if edits else "khác dấu câu hoặc khoảng trắng"
    return (
        f"Có sai khác nội dung so với ground truth Nhóm 1: {edit_summary}. "
        f"Edit distance={distance}; accuracy bỏ xuống dòng={score:.2%}; "
        f"giữ xuống dòng={line_score:.2%}. Cần mở hyperlink ảnh để xác minh "
        "từng chữ, dị thể, thứ tự dòng và lạc khoản trước khi kết luận."
    )


def copy_template_style(template: Path, target) -> None:
    if not template.is_file():
        return
    workbook = load_workbook(template)
    source = workbook.active
    for column_index in range(1, len(OUTPUT_HEADERS) + 1):
        letter = get_column_letter(column_index)
        target.column_dimensions[letter].width = source.column_dimensions[letter].width
        target.cell(1, column_index)._style = copy(source.cell(1, column_index)._style)
        if source.max_row >= 2:
            target.cell(2, column_index)._style = copy(source.cell(2, column_index)._style)
    target.freeze_panes = source.freeze_panes
    target.auto_filter.ref = f"A1:H1"


def main() -> int:
    args = parse_args()
    for path in (args.input, args.ground_truth):
        if not path.is_file():
            raise FileNotFoundError(path)

    truth = load_ground_truth(args.ground_truth)
    # Cột # của workbook nguồn là công thức =ROW()-2; dùng giá trị cache để
    # chọn đúng ID nhưng openpyxl vẫn giữ được hyperlink của ô Image.
    source_book = load_workbook(args.input, data_only=True)
    source_sheet = source_book.active
    columns = header_map(source_sheet)

    selected: list[tuple[int, object]] = []
    seen_ids: set[int] = set()
    for row_index in range(2, source_sheet.max_row + 1):
        raw_id = source_sheet.cell(row_index, columns["#"]).value
        try:
            record_id = int(raw_id)
        except (TypeError, ValueError):
            continue
        if args.start_id <= record_id <= args.end_id:
            if record_id in seen_ids:
                raise ValueError(f"ID bị trùng trong input: {record_id}")
            seen_ids.add(record_id)
            selected.append((record_id, row_index))

    expected = set(range(args.start_id, args.end_id + 1))
    if seen_ids != expected:
        missing = sorted(expected - seen_ids)
        raise ValueError(f"Input thiếu ID: {missing[:10]}")

    output_book = Workbook()
    output_sheet = output_book.active
    output_sheet.title = "Cross-validation"
    copy_template_style(args.template, output_sheet)
    for column_index, header in enumerate(OUTPUT_HEADERS, 1):
        output_sheet.cell(1, column_index, header)

    missing_images: list[str] = []
    unequal = 0
    for output_row, (record_id, source_row) in enumerate(sorted(selected), 2):
        image_cell = source_sheet.cell(source_row, columns["Image"])
        image_name = Path(str(image_cell.value or "")).name
        group1_text = truth.get(image_name)
        if group1_text is None:
            missing_images.append(image_name)
            continue

        group2_text = str(source_sheet.cell(source_row, columns["Label"]).value or "")
        caption = str(source_sheet.cell(source_row, columns["FB Caption"]).value or "")
        score = accuracy(group1_text, group2_text, include_line_breaks=False)
        line_score = accuracy(group1_text, group2_text, include_line_breaks=True)
        unequal += group2_text != group1_text

        hyperlink = image_cell.hyperlink.target if image_cell.hyperlink else image_name
        values = (
            record_id,
            hyperlink,
            group2_text,
            caption,
            group1_text,
            score,
            line_score,
            note_for(group2_text, group1_text, score, line_score),
        )
        for column_index, value in enumerate(values, 1):
            cell = output_sheet.cell(output_row, column_index, value)
            if output_row > 2 and output_sheet.cell(2, column_index).has_style:
                cell._style = copy(output_sheet.cell(2, column_index)._style)
            cell.alignment = copy(cell.alignment)
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="top",
                wrap_text=True,
            )
        output_sheet.cell(output_row, 2).hyperlink = hyperlink
        output_sheet.cell(output_row, 2).style = "Hyperlink"
        output_sheet.cell(output_row, 6).number_format = "0.00%"
        output_sheet.cell(output_row, 7).number_format = "0.00%"
        if group2_text != group1_text:
            output_sheet.cell(output_row, 8).fill = PatternFill("solid", fgColor="FFF2CC")

    if missing_images:
        raise ValueError(f"Không tìm thấy ground truth cho ảnh: {missing_images[:10]}")

    output_sheet.freeze_panes = "A2"
    output_sheet.auto_filter.ref = f"A1:H{output_sheet.max_row}"
    output_sheet.row_dimensions[1].height = 30
    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary = args.output.with_suffix(args.output.suffix + ".tmp")
    output_book.save(temporary)
    os.replace(temporary, args.output)
    print(f"Đã ghi {len(selected)} record (ID {args.start_id}..{args.end_id}) → {args.output}")
    print(f"Khớp tuyệt đối: {len(selected) - unequal}; cần rà soát thủ công: {unequal}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
